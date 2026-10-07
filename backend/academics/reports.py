"""Session reports and exports. Authorization is enforced by the viewset."""
import re
from html import escape
from io import BytesIO
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from accounts.models import StudentProfile
from .models import AttendanceSession, Enrollment
from .serializers import AttendanceSessionSerializer


def session_report(session, request):
    finished = session.start_time <= timezone.now() and session.effective_status in {"closed", "expired"}
    records = {record.student_id: record for record in session.records.select_related("student__user")}
    roster = StudentProfile.objects.filter(
        Q(enrollments__class_group=session.class_group, enrollments__status=Enrollment.Status.ACTIVE,
          enrollments__enrolled_at__lte=session.end_time) | Q(attendance_records__session=session)
    ).select_related("user").distinct().order_by("student_number")
    students = []
    counts = {"present":0, "late":0, "absent":0, "excused":0, "pending":0}
    for student in roster:
        record = records.get(student.pk)
        status = record.status if record else "absent" if finished else "pending"
        counts[status] += 1
        students.append({
            "id":str(student.user_id), "studentId":student.student_number, "name":student.user.full_name,
            "email":student.user.email, "department":student.department, "level":student.level,
            "status":status, "checkInTime":record.check_in_time.isoformat() if record else None,
            "distance":float(record.distance_from_venue) if record else None,
            "accuracy":float(record.gps_accuracy) if record else None,
        })
    return {
        "session":AttendanceSessionSerializer(session, context={"request":request}).data,
        "canExport":finished, "summary":dict(counts, total=len(students)), "students":students,
    }


def export_report(report, file_type):
    session = report["session"]
    code = re.sub(r"[^A-Za-z0-9_-]", "_", session["courseCode"])
    filename = f'{code}-{session["id"]}-attendance'
    headers = ["Student ID", "Name", "Email", "Level", "Department", "Status", "Check-in (UTC)", "Distance (m)", "GPS accuracy (m)"]
    rows = [[s["studentId"], s["name"], s["email"], s["level"], s["department"],
             s["status"].title(), s["checkInTime"] or "", s["distance"], s["accuracy"]] for s in report["students"]]
    output = BytesIO()
    if file_type == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Attendance"
        sheet.append([f'{session["courseCode"]} - {session["courseName"]}'])
        sheet.append(["Class", session["classGroupName"], "Venue", session["venueName"]])
        sheet.append(["Lecturer", session["lecturerName"]])
        sheet.append(["Start (UTC)", session["startTime"], "End (UTC)", session["endTime"]])
        sheet.append(["Total", report["summary"]["total"], "Present", report["summary"]["present"], "Late", report["summary"]["late"], "Absent", report["summary"]["absent"]])
        sheet.append([])
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        # User-supplied text must never become a spreadsheet formula.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in sheet[7]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="235347")
        for column, width in zip("ABCDEFGHI", [22, 30, 38, 12, 30, 15, 32, 18, 20]):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A8"
        sheet.auto_filter.ref = f"A7:I{max(7, sheet.max_row)}"
        workbook.save(output)
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
        styles = getSampleStyleSheet()
        small = ParagraphStyle("ReportCell", fontName="Helvetica", fontSize=8, leading=11)
        story = [Paragraph(escape(f'{session["courseCode"]} - {session["courseName"]}'), styles["Heading1"])]
        for text in [
            f'Class: {session["classGroupName"]} | Venue: {session["venueName"]}',
            f'Lecturer: {session["lecturerName"]}',
            f'Start (UTC): {session["startTime"]} | End (UTC): {session["endTime"]}',
            f'Total: {report["summary"]["total"]} | Present: {report["summary"]["present"]} | Late: {report["summary"]["late"]} | Absent: {report["summary"]["absent"]} | Excused: {report["summary"]["excused"]}',
        ]:
            story.append(Paragraph(escape(text), styles["Normal"]))
            story.append(Spacer(1, 8))
        table_data = [[Paragraph(escape(str(value or "")), small) for value in headers[:7]]]
        table_data.extend([[Paragraph(escape(str(value or "")), small) for value in row[:7]] for row in rows])
        table = Table(table_data, repeatRows=1, colWidths=[85, 115, 165, 40, 125, 60, 150])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbeae4")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f6f4")]),
            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        story.append(table)
        document.build(story)
        content_type = "application/pdf"
    response = HttpResponse(output.getvalue(), content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}.{file_type}"'
    return response
